"""
Module xuất file Microsoft Word (.docx) chuẩn hành chính sư phạm của Bộ GD&ĐT.
Bao gồm:
1. Tiêu đề trường/sở, thông tin kỳ kiểm tra.
2. BẢNG MA TRẬN ĐỀ KIỂM TRA (chuẩn bảng Mẫu 1 hoặc Mẫu 2, có viền, gộp ô, shading xám).
3. BẢNG ĐẶC TẢ KỸ THUẬT ĐỀ KIỂM TRA (đầy đủ Yêu cầu cần đạt 3 mức độ nhận thức).
4. ĐỀ KIỂM TRA ĐỊNH KỲ CHÍNH THỨC (Phần I, II, III, IV).
5. HƯỚNG DẪN CHẤM VÀ ĐÁP ÁN CHI TIẾT (Bảng ma trận điểm, đáp án trắc nghiệm, thang điểm tự luận).
"""

import io
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    """Đặt màu nền cho ô trong bảng (ví dụ: 'F2F2F2' cho xám nhạt).
    Đã kiểm tra hợp lệ chuỗi hex màu để chống tiêm mã XML.
    """
    clean_hex = str(fill_hex).strip().lstrip("#").upper()
    if not re.match(r"^[0-9A-F]{6}$", clean_hex):
        clean_hex = "F2F2F2"
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{clean_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Đặt đệm (padding) trong ô."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_margins(table, top=100, bottom=100, left=140, right=140):
    """Đặt đệm (padding) cho toàn bộ bảng."""
    tblPr = table._tbl.tblPr
    tblCellMar = OxmlElement('w:tblCellMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tblCellMar.append(node)
    tblPr.append(tblCellMar)

def set_table_borders(table, color="B0B0B0", sz="4", val="single"):
    """Đặt viền thanh lịch chuẩn form hành chính cho bảng."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def generate_exam_docx(blueprint: dict, exam_content: dict, school_name: str = "TRƯỜNG THCS TIÊN SƠN", time_minutes: int = 45) -> io.BytesIO:
    # Chuẩn hóa đầu vào người dùng
    school_name = (school_name or "TRƯỜNG THCS TIÊN SƠN").strip()
    if len(school_name) > 120:
        school_name = school_name[:120].strip() + "..."
        
    try:
        time_minutes = int(time_minutes)
        if time_minutes <= 0 or time_minutes > 180:
            time_minutes = 45
    except (ValueError, TypeError):
        time_minutes = 45

    doc = Document()
    
    # Cấu hình font chuẩn Times New Roman cho toàn bộ tài liệu
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(11)
    rFonts = style.element.rPr.get_or_add_rFonts()
    rFonts.set(qn('w:ascii'), 'Times New Roman')
    rFonts.set(qn('w:hAnsi'), 'Times New Roman')
    rFonts.set(qn('w:cs'), 'Times New Roman')
    rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    
    # Phần 1 (Ma trận & Đặc tả): Đặt Khổ ngang (Landscape - A4: 11.69 x 8.27 in) để bảng 19 cột thoáng rộng
    sec_matrix = doc.sections[0]
    sec_matrix.orientation = WD_ORIENT.LANDSCAPE
    sec_matrix.page_width = Inches(11.69)
    sec_matrix.page_height = Inches(8.27)
    sec_matrix.top_margin = Inches(0.55)
    sec_matrix.bottom_margin = Inches(0.55)
    sec_matrix.left_margin = Inches(0.65)
    sec_matrix.right_margin = Inches(0.65)
        
    grade_title = blueprint["grade_title"].upper()
    exam_period = blueprint["exam_period"].upper()
    template_type = blueprint["template_type"]
    summary = blueprint["summary"]
    matrix_rows = blueprint["matrix_rows"]

    dept_name = "UBND PHƯỜNG TIÊN SƠN" if "TIÊN SƠN" in school_name.upper() else "PHÒNG GD&ĐT QUẬN/HUYỆN"

    # ------------------ PHẦN HEADER ĐẦU TRANG ------------------
    header_tbl = doc.add_table(rows=2, cols=2)
    header_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_tbl.autofit = False

    # Cột trái
    cell_tl = header_tbl.cell(0, 0)
    p_tl1 = cell_tl.paragraphs[0]
    p_tl1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p_tl1.add_run(f"{dept_name}\n{school_name.upper()}")
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(11)
    r1.font.bold = True

    p_tl2 = cell_tl.add_paragraph()
    p_tl2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p_tl2.add_run("TỔ CHUYÊN MÔN: TIN HỌC")
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(11)

    # Cột phải
    cell_tr = header_tbl.cell(0, 1)
    p_tr1 = cell_tr.paragraphs[0]
    p_tr1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r3 = p_tr1.add_run(f"HỒ SƠ KHẢO THÍ VÀ KIỂM TRA ĐỊNH KỲ\nBÀI KIỂM TRA: {exam_period}\nNĂM HỌC 2026 - 2027")
    r3.font.name = "Times New Roman"
    r3.font.size = Pt(11)
    r3.font.bold = True

    p_tr2 = cell_tr.add_paragraph()
    p_tr2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r4 = p_tr2.add_run(f"Môn: {grade_title} - Thời gian làm bài: {time_minutes} phút")
    r4.font.name = "Times New Roman"
    r4.font.size = Pt(10)
    r4.font.italic = True

    # Xóa viền header table
    set_table_borders(header_tbl, color="FFFFFF", val="none")

    doc.add_paragraph() # Khoảng cách

    # ==============================================================
    # 1. BẢNG MA TRẬN ĐỀ KIỂM TRA
    # ==============================================================
    p_title1 = doc.add_paragraph()
    r_t1 = p_title1.add_run("I. MA TRẬN ĐỀ KIỂM TRA ĐỊNH KỲ (CHUẨN GDPT 2018)")
    r_t1.font.name = "Times New Roman"
    r_t1.font.size = Pt(13)
    r_t1.font.bold = True
    p_title1.alignment = WD_ALIGN_PARAGRAPH.LEFT

    p_sub1 = doc.add_paragraph()
    r_s1 = p_sub1.add_run(f"* Tỉ lệ phân bố mức độ nhận thức: 40% Nhận biết – 30% Thông hiểu – 30% Vận dụng (Tỉ lệ 4 : 3 : 3)\n* Cấu trúc đề: {summary['ten_mau']}")
    r_s1.font.name = "Times New Roman"
    r_s1.font.size = Pt(10.5)
    r_s1.font.italic = True

    if template_type == "mau_2":
        # MẪU 2: Bảng Ma trận chuẩn tập huấn TS. Vũ Cẩm Tú (19 cột)
        cols_count = 19
        mat_tbl = doc.add_table(rows=3 + len(matrix_rows) + 3, cols=cols_count)
        set_table_borders(mat_tbl)
        set_table_margins(mat_tbl, 60, 60, 70, 70)
        mat_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

        # Đặt tiêu đề Header Row 0
        mat_tbl.cell(0, 0).text = "TT"
        mat_tbl.cell(0, 1).text = "Chủ đề / Chương"
        mat_tbl.cell(0, 2).text = "Nội dung / Đơn vị kiến thức"
        mat_tbl.cell(0, 3).text = "Mức độ đánh giá"
        mat_tbl.cell(0, 15).text = "Tổng"
        mat_tbl.cell(0, 18).text = "Tỉ lệ % điểm"

        # Đặt tiêu đề Header Row 1
        mat_tbl.cell(1, 3).text = "TNKQ (Nhiều lựa chọn)"
        mat_tbl.cell(1, 6).text = "TNKQ (Đúng - Sai)"
        mat_tbl.cell(1, 9).text = "TNKQ (Trả lời ngắn)"
        mat_tbl.cell(1, 12).text = "Tự luận"
        mat_tbl.cell(1, 15).text = "Biết"
        mat_tbl.cell(1, 16).text = "Hiểu"
        mat_tbl.cell(1, 17).text = "Vận dụng"

        # Đặt tiêu đề Header Row 2
        for start_c in [3, 6, 9, 12]:
            mat_tbl.cell(2, start_c).text = "Biết"
            mat_tbl.cell(2, start_c + 1).text = "Hiểu"
            mat_tbl.cell(2, start_c + 2).text = "V.Dụng"

        # Gộp ô (Merge) Header chuẩn xác
        mat_tbl.cell(0, 0).merge(mat_tbl.cell(2, 0))
        mat_tbl.cell(0, 1).merge(mat_tbl.cell(2, 1))
        mat_tbl.cell(0, 2).merge(mat_tbl.cell(2, 2))
        mat_tbl.cell(0, 3).merge(mat_tbl.cell(0, 14))
        mat_tbl.cell(0, 15).merge(mat_tbl.cell(0, 17))
        mat_tbl.cell(0, 18).merge(mat_tbl.cell(2, 18))

        mat_tbl.cell(1, 3).merge(mat_tbl.cell(1, 5))
        mat_tbl.cell(1, 6).merge(mat_tbl.cell(1, 8))
        mat_tbl.cell(1, 9).merge(mat_tbl.cell(1, 11))
        mat_tbl.cell(1, 12).merge(mat_tbl.cell(1, 14))
        mat_tbl.cell(1, 15).merge(mat_tbl.cell(2, 15))
        mat_tbl.cell(1, 16).merge(mat_tbl.cell(2, 16))
        mat_tbl.cell(1, 17).merge(mat_tbl.cell(2, 17))

        # Format header styling
        for r_idx in range(3):
            for c in mat_tbl.rows[r_idx].cells:
                set_cell_background(c, "EAECEE")
                c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                for p in c.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for r in p.runs:
                        r.font.name = "Times New Roman"
                        r.font.size = Pt(8.5)
                        r.font.bold = True

        # Ghi dữ liệu từng dòng
        current_row = 3
        for item in matrix_rows:
            r_cells = mat_tbl.rows[current_row].cells
            r_cells[0].text = str(item["tt"])
            r_cells[1].text = item["chu_de"].split(":")[0]
            r_cells[2].text = item["don_vi"]
            r_cells[3].text = str(item["tn_biet"]) if item["tn_biet"] else ""
            r_cells[4].text = str(item["tn_hieu"]) if item["tn_hieu"] else ""
            r_cells[5].text = ""
            r_cells[6].text = str(item["ds_biet"]) if item["ds_biet"] else ""
            r_cells[7].text = str(item["ds_hieu"]) if item["ds_hieu"] else ""
            r_cells[8].text = ""
            r_cells[9].text = str(item["tln_biet"]) if item["tln_biet"] else ""
            r_cells[10].text = str(item["tln_hieu"]) if item["tln_hieu"] else ""
            r_cells[11].text = ""
            r_cells[12].text = ""
            r_cells[13].text = ""
            r_cells[14].text = str(item["tl_vd"]) if item["tl_vd"] else ""
            r_cells[15].text = str(item["tong_biet"]) if item["tong_biet"] else ""
            r_cells[16].text = str(item["tong_hieu"]) if item["tong_hieu"] else ""
            r_cells[17].text = str(item["tong_vd"]) if item["tong_vd"] else ""
            r_cells[18].text = f"{item['ti_le']}%"

            for c_idx, c in enumerate(r_cells):
                c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                for p in c.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx in [1, 2] else WD_ALIGN_PARAGRAPH.CENTER
                    for r in p.runs:
                        r.font.name = "Times New Roman"
                        r.font.size = Pt(8.5)
            current_row += 1

        # Dòng Tổng số câu (Tính động từ các dòng)
        r_sc = current_row
        r_sd = current_row + 1
        r_tl = current_row + 2

        tot_tn_b = sum(r["tn_biet"] for r in matrix_rows)
        tot_tn_h = sum(r["tn_hieu"] for r in matrix_rows)
        tot_ds_b = sum(r["ds_biet"] for r in matrix_rows)
        tot_ds_h = sum(r["ds_hieu"] for r in matrix_rows)
        tot_tln_b = sum(r["tln_biet"] for r in matrix_rows)
        tot_tln_h = sum(r["tln_hieu"] for r in matrix_rows)
        tot_tl_v = sum(r["tl_vd"] for r in matrix_rows)
        tot_b = sum(r["tong_biet"] for r in matrix_rows)
        tot_h = sum(r["tong_hieu"] for r in matrix_rows)
        tot_v = sum(r["tong_vd"] for r in matrix_rows)
        tot_cau = sum(r["tong_cau"] for r in matrix_rows)

        mat_tbl.cell(r_sc, 0).text = "Tổng số câu"
        mat_tbl.cell(r_sc, 3).text = str(tot_tn_b)
        mat_tbl.cell(r_sc, 4).text = str(tot_tn_h)
        mat_tbl.cell(r_sc, 6).text = str(tot_ds_b)
        mat_tbl.cell(r_sc, 7).text = str(tot_ds_h)
        mat_tbl.cell(r_sc, 9).text = str(tot_tln_b)
        mat_tbl.cell(r_sc, 10).text = str(tot_tln_h)
        mat_tbl.cell(r_sc, 14).text = str(tot_tl_v)
        mat_tbl.cell(r_sc, 15).text = str(tot_b)
        mat_tbl.cell(r_sc, 16).text = str(tot_h)
        mat_tbl.cell(r_sc, 17).text = str(tot_v)
        mat_tbl.cell(r_sc, 18).text = f"{tot_cau} câu"

        # Dòng Tổng số điểm
        mat_tbl.cell(r_sd, 0).text = "Tổng số điểm"
        mat_tbl.cell(r_sd, 3).text = "3,0 điểm"
        mat_tbl.cell(r_sd, 6).text = "2,0 điểm"
        mat_tbl.cell(r_sd, 9).text = "2,0 điểm"
        mat_tbl.cell(r_sd, 12).text = "3,0 điểm"
        mat_tbl.cell(r_sd, 15).text = "4,0 điểm"
        mat_tbl.cell(r_sd, 16).text = "3,0 điểm"
        mat_tbl.cell(r_sd, 17).text = "3,0 điểm"
        mat_tbl.cell(r_sd, 18).text = "10,0 điểm"

        # Dòng Tỉ lệ %
        mat_tbl.cell(r_tl, 0).text = "Tỉ lệ % điểm"
        mat_tbl.cell(r_tl, 3).text = "30%"
        mat_tbl.cell(r_tl, 6).text = "20%"
        mat_tbl.cell(r_tl, 9).text = "20%"
        mat_tbl.cell(r_tl, 12).text = "30%"
        mat_tbl.cell(r_tl, 15).text = "40%"
        mat_tbl.cell(r_tl, 16).text = "30%"
        mat_tbl.cell(r_tl, 17).text = "30%"
        mat_tbl.cell(r_tl, 18).text = "100%"

        # Gộp ô dòng summary
        mat_tbl.cell(r_sc, 0).merge(mat_tbl.cell(r_sc, 2))
        mat_tbl.cell(r_sd, 0).merge(mat_tbl.cell(r_sd, 2))
        mat_tbl.cell(r_tl, 0).merge(mat_tbl.cell(r_tl, 2))

        mat_tbl.cell(r_sd, 3).merge(mat_tbl.cell(r_sd, 5))
        mat_tbl.cell(r_sd, 6).merge(mat_tbl.cell(r_sd, 8))
        mat_tbl.cell(r_sd, 9).merge(mat_tbl.cell(r_sd, 11))
        mat_tbl.cell(r_sd, 12).merge(mat_tbl.cell(r_sd, 14))

        mat_tbl.cell(r_tl, 3).merge(mat_tbl.cell(r_tl, 5))
        mat_tbl.cell(r_tl, 6).merge(mat_tbl.cell(r_tl, 8))
        mat_tbl.cell(r_tl, 9).merge(mat_tbl.cell(r_tl, 11))
        mat_tbl.cell(r_tl, 12).merge(mat_tbl.cell(r_tl, 14))

        for r_offset in range(3):
            for c in mat_tbl.rows[current_row + r_offset].cells:
                set_cell_background(c, "F8F9FA")
                c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                for p in c.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for r in p.runs:
                        r.font.name = "Times New Roman"
                        r.font.size = Pt(8.5)
                        r.font.bold = True

    else:
        # MẪU 1: Bảng Ma trận chuẩn docx trường (16 cột)
        cols_count = 16
        mat_tbl = doc.add_table(rows=4 + len(matrix_rows) + 3, cols=cols_count)
        set_table_borders(mat_tbl)
        set_table_margins(mat_tbl, 60, 60, 70, 70)
        mat_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

        # Row 0
        mat_tbl.cell(0, 0).text = "TT"
        mat_tbl.cell(0, 1).text = "Chương / Chủ đề"
        mat_tbl.cell(0, 2).text = "Nội dung / Đơn vị kiến thức"
        mat_tbl.cell(0, 3).text = "Số câu hỏi / ý hỏi ở các mức độ đánh giá"
        mat_tbl.cell(0, 12).text = "Tổng"
        mat_tbl.cell(0, 15).text = "Tỉ lệ % điểm"

        # Row 1
        mat_tbl.cell(1, 3).text = "Trắc nghiệm khách quan"
        mat_tbl.cell(1, 9).text = "Tự luận"
        mat_tbl.cell(1, 12).text = "Biết"
        mat_tbl.cell(1, 13).text = "Hiểu"
        mat_tbl.cell(1, 14).text = "Vận dụng"

        # Row 2
        mat_tbl.cell(2, 3).text = "Nhiều lựa chọn"
        mat_tbl.cell(2, 6).text = "Đúng / Sai"

        # Row 3
        for sc in [3, 6, 9]:
            mat_tbl.cell(3, sc).text = "Biết"
            mat_tbl.cell(3, sc + 1).text = "Hiểu"
            mat_tbl.cell(3, sc + 2).text = "V.Dụng"

        # Merges cho Mẫu 1
        mat_tbl.cell(0, 0).merge(mat_tbl.cell(3, 0))
        mat_tbl.cell(0, 1).merge(mat_tbl.cell(3, 1))
        mat_tbl.cell(0, 2).merge(mat_tbl.cell(3, 2))
        mat_tbl.cell(0, 3).merge(mat_tbl.cell(0, 11))
        mat_tbl.cell(0, 12).merge(mat_tbl.cell(0, 14))
        mat_tbl.cell(0, 15).merge(mat_tbl.cell(3, 15))

        mat_tbl.cell(1, 3).merge(mat_tbl.cell(1, 8))
        mat_tbl.cell(1, 9).merge(mat_tbl.cell(2, 11))

        mat_tbl.cell(1, 12).merge(mat_tbl.cell(3, 12))
        mat_tbl.cell(1, 13).merge(mat_tbl.cell(3, 13))
        mat_tbl.cell(1, 14).merge(mat_tbl.cell(3, 14))

        mat_tbl.cell(2, 3).merge(mat_tbl.cell(2, 5))
        mat_tbl.cell(2, 6).merge(mat_tbl.cell(2, 8))

        for r_idx in range(4):
            for c in mat_tbl.rows[r_idx].cells:
                set_cell_background(c, "EAECEE")
                c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                for p in c.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for r in p.runs:
                        r.font.name = "Times New Roman"
                        r.font.size = Pt(8.5)
                        r.font.bold = True

        current_row = 4
        for item in matrix_rows:
            r_cells = mat_tbl.rows[current_row].cells
            r_cells[0].text = str(item["tt"])
            r_cells[1].text = item["chu_de"].split(":")[0]
            r_cells[2].text = item["don_vi"]
            r_cells[3].text = str(item["tn_biet"]) if item["tn_biet"] else ""
            r_cells[4].text = str(item["tn_hieu"]) if item["tn_hieu"] else ""
            r_cells[5].text = ""
            r_cells[6].text = str(item["ds_biet"]) if item["ds_biet"] else ""
            r_cells[7].text = str(item["ds_hieu"]) if item["ds_hieu"] else ""
            r_cells[8].text = ""
            r_cells[9].text = str(item.get("tl_biet", 0)) if item.get("tl_biet") else ""
            r_cells[10].text = str(item.get("tl_hieu", 0)) if item.get("tl_hieu") else ""
            r_cells[11].text = str(item.get("tl_vd", 0)) if item.get("tl_vd") else ""
            r_cells[12].text = str(item["tong_biet"]) if item["tong_biet"] else ""
            r_cells[13].text = str(item["tong_hieu"]) if item["tong_hieu"] else ""
            r_cells[14].text = str(item["tong_vd"]) if item["tong_vd"] else ""
            r_cells[15].text = f"{item['ti_le']}%"

            for c_idx, c in enumerate(r_cells):
                c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                for p in c.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx in [1, 2] else WD_ALIGN_PARAGRAPH.CENTER
                    for r in p.runs:
                        r.font.name = "Times New Roman"
                        r.font.size = Pt(8.5)
            current_row += 1

        # Summary Mẫu 1
        r_sc = current_row
        r_sd = current_row + 1
        r_tl = current_row + 2

        tot_tn_b = sum(r["tn_biet"] for r in matrix_rows)
        tot_tn_h = sum(r["tn_hieu"] for r in matrix_rows)
        tot_ds_b = sum(r["ds_biet"] for r in matrix_rows)
        tot_ds_h = sum(r["ds_hieu"] for r in matrix_rows)
        tot_tl_b = sum(r.get("tl_biet", 0) for r in matrix_rows)
        tot_tl_h = sum(r.get("tl_hieu", 0) for r in matrix_rows)
        tot_tl_v = sum(r.get("tl_vd", 0) for r in matrix_rows)
        tot_b = sum(r["tong_biet"] for r in matrix_rows)
        tot_h = sum(r["tong_hieu"] for r in matrix_rows)
        tot_v = sum(r["tong_vd"] for r in matrix_rows)
        tot_cau = sum(r["tong_cau"] for r in matrix_rows)

        mat_tbl.cell(r_sc, 0).text = "Tổng số câu"
        mat_tbl.cell(r_sc, 3).text = str(tot_tn_b)
        mat_tbl.cell(r_sc, 4).text = str(tot_tn_h)
        mat_tbl.cell(r_sc, 6).text = str(tot_ds_b)
        mat_tbl.cell(r_sc, 7).text = str(tot_ds_h)
        mat_tbl.cell(r_sc, 9).text = str(tot_tl_b)
        mat_tbl.cell(r_sc, 10).text = str(tot_tl_h)
        mat_tbl.cell(r_sc, 11).text = str(tot_tl_v)
        mat_tbl.cell(r_sc, 12).text = str(tot_b)
        mat_tbl.cell(r_sc, 13).text = str(tot_h)
        mat_tbl.cell(r_sc, 14).text = str(tot_v)
        mat_tbl.cell(r_sc, 15).text = f"{tot_cau} câu"

        mat_tbl.cell(r_sd, 0).text = "Tổng số điểm"
        mat_tbl.cell(r_sd, 3).text = "6,0 điểm"
        mat_tbl.cell(r_sd, 6).text = "1,0 điểm"
        mat_tbl.cell(r_sd, 9).text = "3,0 điểm"
        mat_tbl.cell(r_sd, 12).text = "4,0 điểm"
        mat_tbl.cell(r_sd, 13).text = "3,0 điểm"
        mat_tbl.cell(r_sd, 14).text = "3,0 điểm"
        mat_tbl.cell(r_sd, 15).text = "10,0 điểm"

        mat_tbl.cell(r_tl, 0).text = "Tỉ lệ %"
        mat_tbl.cell(r_tl, 3).text = "60%"
        mat_tbl.cell(r_tl, 6).text = "10%"
        mat_tbl.cell(r_tl, 9).text = "30%"
        mat_tbl.cell(r_tl, 12).text = "40%"
        mat_tbl.cell(r_tl, 13).text = "30%"
        mat_tbl.cell(r_tl, 14).text = "30%"
        mat_tbl.cell(r_tl, 15).text = "100%"

        # Gộp ô summary Mẫu 1
        mat_tbl.cell(r_sc, 0).merge(mat_tbl.cell(r_sc, 2))
        mat_tbl.cell(r_sd, 0).merge(mat_tbl.cell(r_sd, 2))
        mat_tbl.cell(r_tl, 0).merge(mat_tbl.cell(r_tl, 2))

        mat_tbl.cell(r_sd, 3).merge(mat_tbl.cell(r_sd, 5))
        mat_tbl.cell(r_sd, 6).merge(mat_tbl.cell(r_sd, 8))
        mat_tbl.cell(r_sd, 9).merge(mat_tbl.cell(r_sd, 11))

        mat_tbl.cell(r_tl, 3).merge(mat_tbl.cell(r_tl, 5))
        mat_tbl.cell(r_tl, 6).merge(mat_tbl.cell(r_tl, 8))
        mat_tbl.cell(r_tl, 9).merge(mat_tbl.cell(r_tl, 11))

        for r_offset in range(3):
            for c in mat_tbl.rows[current_row + r_offset].cells:
                set_cell_background(c, "F8F9FA")
                c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                for p in c.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for r in p.runs:
                        r.font.name = "Times New Roman"
                        r.font.size = Pt(8.5)
                        r.font.bold = True

    doc.add_page_break()

    # ==============================================================
    # 2. BẢNG ĐẶC TẢ ĐỀ KIỂM TRA
    # ==============================================================
    p_title2 = doc.add_paragraph()
    r_t2 = p_title2.add_run("II. BẢN ĐẶC TẢ KỸ THUẬT ĐỀ KIỂM TRA ĐỊNH KỲ")
    r_t2.font.name = "Times New Roman"
    r_t2.font.size = Pt(13)
    r_t2.font.bold = True

    if template_type == "mau_1":
        # BẢN ĐẶC TẢ MẪU 1 CHUẨN TRƯỜNG THCS (13 cột, tách 3 mức độ Nhận biết / Thông hiểu / Vận dụng cho từng bài học)
        cols_count = 13
        total_spec_rows = 4 + (len(matrix_rows) * 3) + 3
        spec_tbl = doc.add_table(rows=total_spec_rows, cols=cols_count)
        set_table_borders(spec_tbl)
        set_table_margins(spec_tbl, 50, 50, 60, 60)
        spec_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

        # Hàng 0
        spec_tbl.cell(0, 0).text = "TT"
        spec_tbl.cell(0, 1).text = "Chương / chủ đề"
        spec_tbl.cell(0, 2).text = "Nội dung / đơn vị kiến thức"
        spec_tbl.cell(0, 3).text = "Yêu cầu cần đạt\n(theo 3 mức độ)"
        spec_tbl.cell(0, 4).text = "Số câu hỏi / ý hỏi ở các mức độ đánh giá"

        # Hàng 1
        spec_tbl.cell(1, 4).text = "Trắc nghiệm khách quan"
        spec_tbl.cell(1, 10).text = "Tự luận"

        # Hàng 2
        spec_tbl.cell(2, 4).text = "Nhiều lựa chọn"
        spec_tbl.cell(2, 7).text = "Đúng / Sai"
        spec_tbl.cell(2, 10).text = "Tự luận"

        # Hàng 3
        for sc in [4, 7, 10]:
            spec_tbl.cell(3, sc).text = "Biết"
            spec_tbl.cell(3, sc + 1).text = "Hiểu"
            spec_tbl.cell(3, sc + 2).text = "V.Dụng"

        # Gộp ô Header
        spec_tbl.cell(0, 0).merge(spec_tbl.cell(3, 0))
        spec_tbl.cell(0, 1).merge(spec_tbl.cell(3, 1))
        spec_tbl.cell(0, 2).merge(spec_tbl.cell(3, 2))
        spec_tbl.cell(0, 3).merge(spec_tbl.cell(3, 3))
        spec_tbl.cell(0, 4).merge(spec_tbl.cell(0, 12))

        spec_tbl.cell(1, 4).merge(spec_tbl.cell(1, 9))
        spec_tbl.cell(1, 10).merge(spec_tbl.cell(2, 12))

        spec_tbl.cell(2, 4).merge(spec_tbl.cell(2, 6))
        spec_tbl.cell(2, 7).merge(spec_tbl.cell(2, 9))

        for r_idx in range(4):
            for c in spec_tbl.rows[r_idx].cells:
                set_cell_background(c, "EAECEE")
                c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                for p in c.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for r in p.runs:
                        r.font.name = "Times New Roman"
                        r.font.size = Pt(8.5)
                        r.font.bold = True

        # Render dữ liệu bài học (mỗi bài 3 dòng Biết, Hiểu, Vận dụng)
        curr_row = 4
        for item in matrix_rows:
            r_b = curr_row
            r_h = curr_row + 1
            r_v = curr_row + 2

            # Gộp cột 0, 1, 2 theo chiều dọc cho bài học
            spec_tbl.cell(r_b, 0).merge(spec_tbl.cell(r_v, 0))
            spec_tbl.cell(r_b, 1).merge(spec_tbl.cell(r_v, 1))
            spec_tbl.cell(r_b, 2).merge(spec_tbl.cell(r_v, 2))

            spec_tbl.cell(r_b, 0).text = str(item["tt"])
            spec_tbl.cell(r_b, 1).text = item["chu_de"].split(":")[0]
            spec_tbl.cell(r_b, 2).text = item["don_vi"]

            # 1. Dòng Nhận biết
            yccd_b = "Nhận biết:\n- " + "\n- ".join(item["yccd"].get("biet", ["Nhận biết khái niệm, quy tắc cơ bản."]))
            spec_tbl.cell(r_b, 3).text = yccd_b
            spec_tbl.cell(r_b, 4).text = str(item["tn_biet"]) if item["tn_biet"] else ""
            spec_tbl.cell(r_b, 5).text = ""
            spec_tbl.cell(r_b, 6).text = ""
            spec_tbl.cell(r_b, 7).text = str(item["ds_biet"]) if item.get("ds_biet") else ""
            spec_tbl.cell(r_b, 8).text = ""
            spec_tbl.cell(r_b, 9).text = ""
            spec_tbl.cell(r_b, 10).text = str(item.get("tl_biet", 0)) if item.get("tl_biet") else ""
            spec_tbl.cell(r_b, 11).text = ""
            spec_tbl.cell(r_b, 12).text = ""

            # 2. Dòng Thông hiểu
            yccd_h = "Thông hiểu:\n- " + "\n- ".join(item["yccd"].get("hieu", ["Giải thích, phân biệt, thao tác mẫu."]))
            spec_tbl.cell(r_h, 3).text = yccd_h
            spec_tbl.cell(r_h, 4).text = ""
            spec_tbl.cell(r_h, 5).text = str(item["tn_hieu"]) if item["tn_hieu"] else ""
            spec_tbl.cell(r_h, 6).text = ""
            spec_tbl.cell(r_h, 7).text = ""
            spec_tbl.cell(r_h, 8).text = str(item["ds_hieu"]) if item.get("ds_hieu") else ""
            spec_tbl.cell(r_h, 9).text = ""
            spec_tbl.cell(r_h, 10).text = ""
            spec_tbl.cell(r_h, 11).text = str(item.get("tl_hieu", 0)) if item.get("tl_hieu") else ""
            spec_tbl.cell(r_h, 12).text = ""

            # 3. Dòng Vận dụng
            yccd_v = "Vận dụng:\n- " + "\n- ".join(item["yccd"].get("van_dung", ["Giải quyết vấn đề thực tiễn."]))
            spec_tbl.cell(r_v, 3).text = yccd_v
            spec_tbl.cell(r_v, 4).text = ""
            spec_tbl.cell(r_v, 5).text = ""
            spec_tbl.cell(r_v, 6).text = ""
            spec_tbl.cell(r_v, 7).text = ""
            spec_tbl.cell(r_v, 8).text = ""
            spec_tbl.cell(r_v, 9).text = ""
            spec_tbl.cell(r_v, 10).text = ""
            spec_tbl.cell(r_v, 11).text = ""
            spec_tbl.cell(r_v, 12).text = str(item.get("tl_vd", 0)) if item.get("tl_vd") else ""

            # Format chữ và căn lề cho 3 dòng này
            for sub_r in [r_b, r_h, r_v]:
                for c_idx, cell in enumerate(spec_tbl.rows[sub_r].cells):
                    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                    for p in cell.paragraphs:
                        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx in [1, 2, 3] else WD_ALIGN_PARAGRAPH.CENTER
                        for r in p.runs:
                            r.font.name = "Times New Roman"
                            r.font.size = Pt(8.5)

            curr_row += 3

        # Dòng Tổng kết đặc tả Mẫu 1
        r_sc = curr_row
        r_sd = curr_row + 1
        r_tl = curr_row + 2

        tot_tn_b = sum(r["tn_biet"] for r in matrix_rows)
        tot_tn_h = sum(r["tn_hieu"] for r in matrix_rows)
        tot_ds_b = sum(r["ds_biet"] for r in matrix_rows)
        tot_ds_h = sum(r["ds_hieu"] for r in matrix_rows)
        tot_tl_v = sum(r.get("tl_vd", 0) for r in matrix_rows)

        spec_tbl.cell(r_sc, 0).text = "Tổng số câu"
        spec_tbl.cell(r_sc, 4).text = str(tot_tn_b)
        spec_tbl.cell(r_sc, 5).text = str(tot_tn_h)
        spec_tbl.cell(r_sc, 6).text = "0"
        spec_tbl.cell(r_sc, 7).text = str(tot_ds_b)
        spec_tbl.cell(r_sc, 8).text = str(tot_ds_h)
        spec_tbl.cell(r_sc, 9).text = "0"
        spec_tbl.cell(r_sc, 10).text = "0"
        spec_tbl.cell(r_sc, 11).text = "0"
        spec_tbl.cell(r_sc, 12).text = str(tot_tl_v)

        spec_tbl.cell(r_sd, 0).text = "Tổng số điểm"
        spec_tbl.cell(r_sd, 4).text = "6,0 điểm"
        spec_tbl.cell(r_sd, 7).text = "1,0 điểm"
        spec_tbl.cell(r_sd, 10).text = "3,0 điểm"

        spec_tbl.cell(r_tl, 0).text = "Tỉ lệ %"
        spec_tbl.cell(r_tl, 4).text = "60%"
        spec_tbl.cell(r_tl, 7).text = "10%"
        spec_tbl.cell(r_tl, 10).text = "30%"

        # Gộp ô summary đặc tả Mẫu 1
        spec_tbl.cell(r_sc, 0).merge(spec_tbl.cell(r_sc, 3))
        spec_tbl.cell(r_sd, 0).merge(spec_tbl.cell(r_sd, 3))
        spec_tbl.cell(r_tl, 0).merge(spec_tbl.cell(r_tl, 3))

        spec_tbl.cell(r_sd, 4).merge(spec_tbl.cell(r_sd, 6))
        spec_tbl.cell(r_sd, 7).merge(spec_tbl.cell(r_sd, 9))
        spec_tbl.cell(r_sd, 10).merge(spec_tbl.cell(r_sd, 12))

        spec_tbl.cell(r_tl, 4).merge(spec_tbl.cell(r_tl, 6))
        spec_tbl.cell(r_tl, 7).merge(spec_tbl.cell(r_tl, 9))
        spec_tbl.cell(r_tl, 10).merge(spec_tbl.cell(r_tl, 12))

        for r_offset in range(3):
            for c in spec_tbl.rows[curr_row + r_offset].cells:
                set_cell_background(c, "F8F9FA")
                c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                for p in c.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for r in p.runs:
                        r.font.name = "Times New Roman"
                        r.font.size = Pt(8.5)
                        r.font.bold = True

        # Thiết lập độ rộng cột chuẩn A4 Landscape cho bảng 13 cột
        col_w_13 = [
            Inches(0.35), Inches(1.1), Inches(1.4), Inches(3.5),
            Inches(0.45), Inches(0.45), Inches(0.45),
            Inches(0.45), Inches(0.45), Inches(0.45),
            Inches(0.45), Inches(0.45), Inches(0.45)
        ]
        for row in spec_tbl.rows:
            for c_idx, w in enumerate(col_w_13):
                row.cells[c_idx].width = w

    else:
        # BẢN ĐẶC TẢ MẪU 2 (Bảng 5 cột tổng hợp theo từng dạng thức)
        spec_tbl = doc.add_table(rows=1 + len(matrix_rows) + 1, cols=5)
        set_table_borders(spec_tbl)
        set_table_margins(spec_tbl, 80, 80, 100, 100)
        spec_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

        spec_headers = ["TT", "Chủ đề / Chương", "Nội dung kiến thức", "Yêu cầu cần đạt (theo chuẩn GDPT 2018)", "Số câu / Mức độ đánh giá"]
        for i, h_text in enumerate(spec_headers):
            c = spec_tbl.cell(0, i)
            c.text = h_text
            set_cell_background(c, "EAECEE")
            for p in c.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.name = "Times New Roman"
                    r.font.size = Pt(9.5)
                    r.font.bold = True

        for idx, item in enumerate(matrix_rows):
            row_c = spec_tbl.rows[idx + 1].cells
            row_c[0].text = str(item["tt"])
            row_c[1].text = item["chu_de"]
            row_c[2].text = item["don_vi"]
            
            # Render YCCĐ chi tiết 3 mức độ
            yccd_text = (
                f"Nhận biết:\n- " + "\n- ".join(item["yccd"].get("biet", ["Khái niệm cơ bản."])) + "\n\n"
                f"Thông hiểu:\n- " + "\n- ".join(item["yccd"].get("hieu", ["Nguyên lý hoạt động."])) + "\n\n"
                f"Vận dụng:\n- " + "\n- ".join(item["yccd"].get("van_dung", ["Áp dụng thực tiễn."]))
            )
            row_c[3].text = yccd_text

            # Số câu hỏi tổng hợp chuẩn xác theo từng dạng
            parts = []
            if item.get("tn_biet") or item.get("tn_hieu"):
                parts.append(f"• TN nhiều lựa chọn: {item['tn_biet']} Biết, {item['tn_hieu']} Hiểu")
            if item.get("ds_biet") or item.get("ds_hieu"):
                parts.append(f"• TN Đúng/Sai: {item['ds_biet']} câu Biết, {item['ds_hieu']} câu Hiểu")
            if item.get("tln_biet") or item.get("tln_hieu"):
                parts.append(f"• TN Trả lời ngắn: {item['tln_biet']} Biết, {item['tln_hieu']} Hiểu")
                
            tl_parts = []
            if item.get("tl_biet"):
                tl_parts.append(f"{item['tl_biet']} Biết")
            if item.get("tl_hieu"):
                tl_parts.append(f"{item['tl_hieu']} Hiểu")
            if item.get("tl_vd"):
                tl_parts.append(f"{item['tl_vd']} Vận dụng")
            if tl_parts:
                parts.append(f"• Tự luận: {', '.join(tl_parts)}")

            row_c[4].text = "\n".join(parts) if parts else "—"

            for c_idx, c in enumerate(row_c):
                for p in c.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx in [1, 2, 3, 4] else WD_ALIGN_PARAGRAPH.CENTER
                    for r in p.runs:
                        r.font.name = "Times New Roman"
                        r.font.size = Pt(9)

        # Dòng tổng kết đặc tả (Gộp ô cột 0..2)
        spec_tbl.cell(-1, 0).merge(spec_tbl.cell(-1, 2))
        spec_tbl.cell(-1, 0).text = "Tổng"
        spec_tbl.cell(-1, 3).text = f"Đảm bảo tỷ lệ đánh giá: 40% Nhận biết ({summary['tong_diem_biet']}đ) – 30% Thông hiểu ({summary['tong_diem_hieu']}đ) – 30% Vận dụng ({summary['tong_diem_vd']}đ)"
        spec_tbl.cell(-1, 4).text = "10.0 điểm\n(100%)"
        for c in [spec_tbl.cell(-1, 0), spec_tbl.cell(-1, 3), spec_tbl.cell(-1, 4)]:
            set_cell_background(c, "F8F9FA")
            c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for p in c.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c != spec_tbl.cell(-1, 3) else WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs:
                    r.font.name = "Times New Roman"
                    r.font.size = Pt(9.5)
                    r.font.bold = True

        # Thiết lập độ rộng cột chuẩn mực cho bản đặc tả 5 cột
        col_widths = [Inches(0.45), Inches(1.5), Inches(1.85), Inches(4.8), Inches(1.8)]
        for row in spec_tbl.rows:
            for c_idx, w in enumerate(col_widths):
                row.cells[c_idx].width = w

    # ==============================================================
    # 3. ĐỀ KIỂM TRA CHÍNH THỨC (Chuyển sang Khổ dọc - Portrait)
    # ==============================================================
    sec_exam = doc.add_section()
    sec_exam.orientation = WD_ORIENT.PORTRAIT
    sec_exam.page_width = Inches(8.27)
    sec_exam.page_height = Inches(11.69)
    sec_exam.top_margin = Inches(0.79)
    sec_exam.bottom_margin = Inches(0.79)
    sec_exam.left_margin = Inches(0.79)
    sec_exam.right_margin = Inches(0.79)

    # ------------------ TIÊU ĐỀ ĐỀ KIỂM TRA CHUẨN 2 CỘT HÀNH CHÍNH ------------------
    ex_head_tbl = doc.add_table(rows=1, cols=2)
    ex_head_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    ex_head_tbl.autofit = False
    set_table_borders(ex_head_tbl, color="FFFFFF", val="none")

    c_left = ex_head_tbl.cell(0, 0)
    c_left.width = Inches(2.9)
    p_l = c_left.paragraphs[0]
    p_l.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_dept = p_l.add_run(f"{dept_name}\n{school_name.upper()}\n")
    r_dept.font.name = "Times New Roman"
    r_dept.font.size = Pt(10.5)
    r_dept.font.bold = True
    r_line = p_l.add_run("-----------------------")
    r_line.font.name = "Times New Roman"
    r_line.font.size = Pt(9)

    c_right = ex_head_tbl.cell(0, 1)
    c_right.width = Inches(3.8)
    p_r = c_right.paragraphs[0]
    p_r.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_ex = p_r.add_run(f"ĐỀ KIỂM TRA {exam_period.upper()}\nNĂM HỌC 2026 - 2027\nMÔN: {grade_title}\n")
    r_ex.font.name = "Times New Roman"
    r_ex.font.size = Pt(11)
    r_ex.font.bold = True
    r_t = p_r.add_run(f"Thời gian làm bài: {time_minutes} phút\n(Không kể thời gian phát đề)")
    r_t.font.name = "Times New Roman"
    r_t.font.size = Pt(9.5)
    r_t.font.italic = True

    # Khung thông tin học sinh
    p_stu = doc.add_paragraph()
    p_stu.paragraph_format.space_before = Pt(8)
    p_stu.paragraph_format.space_after = Pt(4)
    r_stu = p_stu.add_run("Họ và tên học sinh: ................................................................ Lớp: ............. SBD: ..............")
    r_stu.font.name = "Times New Roman"
    r_stu.font.size = Pt(10.5)

    # Khung Điểm và Lời phê giáo viên (Chuẩn giấy thi cấp 2)
    score_tbl = doc.add_table(rows=2, cols=2)
    score_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(score_tbl, color="000000", sz="4", val="single")
    set_table_margins(score_tbl, top=60, bottom=60, left=100, right=100)

    c_s0 = score_tbl.cell(0, 0)
    c_s0.width = Inches(2.2)
    p_s0 = c_s0.paragraphs[0]
    p_s0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_s0 = p_s0.add_run("ĐIỂM")
    r_s0.font.name = "Times New Roman"
    r_s0.font.size = Pt(10)
    r_s0.font.bold = True

    c_s1 = score_tbl.cell(0, 1)
    c_s1.width = Inches(4.5)
    p_s1 = c_s1.paragraphs[0]
    p_s1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_s1 = p_s1.add_run("LỜI PHÊ CỦA THẦY / CÔ GIÁO")
    r_s1.font.name = "Times New Roman"
    r_s1.font.size = Pt(10)
    r_s1.font.bold = True

    # Ô thân để trống chấm điểm
    c_b0 = score_tbl.cell(1, 0)
    c_b0.width = Inches(2.2)
    p_b0 = c_b0.paragraphs[0]
    p_b0.paragraph_format.space_before = Pt(18)
    p_b0.paragraph_format.space_after = Pt(18)

    c_b1 = score_tbl.cell(1, 1)
    c_b1.width = Inches(4.5)
    p_b1 = c_b1.paragraphs[0]
    p_b1.paragraph_format.space_before = Pt(18)
    p_b1.paragraph_format.space_after = Pt(18)

    doc.add_paragraph() # Khoảng cách nhỏ trước Phần I

    # PHẦN I
    p_p1 = doc.add_paragraph()
    r_p1 = p_p1.add_run(f"PHẦN I. Câu hỏi trắc nghiệm nhiều phương án lựa chọn ({summary['diem_tn']} điểm)")
    r_p1.font.name = "Times New Roman"
    r_p1.font.size = Pt(11.5)
    r_p1.font.bold = True
    p_p1_note = doc.add_paragraph("Thí sinh trả lời từ câu 1 đến câu " + str(summary['so_cau_tn']) + ". Mỗi câu hỏi chỉ chọn một phương án đúng duy nhất.")
    p_p1_note.runs[0].font.italic = True
    p_p1_note.runs[0].font.size = Pt(10)

    for item in exam_content.get("phan_I_trac_nghiem", []):
        p_q = doc.add_paragraph()
        r_q = p_q.add_run(f"Câu {item['cau_so']}: ")
        r_q.font.bold = True
        p_q.add_run(item["cau_hoi"])
        
        choices = item.get("lua_chon", {})
        for opt in ["A", "B", "C", "D"]:
            if opt in choices:
                p_opt = doc.add_paragraph()
                p_opt.paragraph_format.left_indent = Inches(0.25)
                r_opt = p_opt.add_run(f"{opt}. ")
                r_opt.font.bold = True
                p_opt.add_run(choices[opt])

    # PHẦN II
    doc.add_paragraph()
    p_p2 = doc.add_paragraph()
    r_p2 = p_p2.add_run(f"PHẦN II. Câu hỏi trắc nghiệm Đúng / Sai ({summary['diem_ds']} điểm)")
    r_p2.font.name = "Times New Roman"
    r_p2.font.size = Pt(11.5)
    r_p2.font.bold = True
    p_p2_note = doc.add_paragraph("Thí sinh trả lời từ câu 1 đến câu " + str(summary['so_cau_ds']) + ". Trong mỗi ý a), b), c), d) ở mỗi câu, thí sinh chọn Đúng hoặc Sai.")
    p_p2_note.runs[0].font.italic = True
    p_p2_note.runs[0].font.size = Pt(10)

    for item in exam_content.get("phan_II_dung_sai", []):
        p_q = doc.add_paragraph()
        r_q = p_q.add_run(f"Câu {item['cau_so']}: ")
        r_q.font.bold = True
        p_q.add_run(item["ngu_canh"])

        for y in item.get("y_hoi", []):
            p_y = doc.add_paragraph()
            p_y.paragraph_format.left_indent = Inches(0.25)
            r_y = p_y.add_run(f"{y['ky_hieu']}) ")
            r_y.font.bold = True
            p_y.add_run(y["noi_dung"])

    # PHẦN III (nếu là Mẫu 2)
    if template_type == "mau_2":
        doc.add_paragraph()
        p_p3 = doc.add_paragraph()
        r_p3 = p_p3.add_run(f"PHẦN III. Câu hỏi trắc nghiệm trả lời ngắn ({summary.get('diem_tln', 2.0)} điểm)")
        r_p3.font.name = "Times New Roman"
        r_p3.font.size = Pt(11.5)
        r_p3.font.bold = True
        p_p3_note = doc.add_paragraph("Thí sinh trả lời từ câu 1 đến câu " + str(summary.get('so_cau_tln', 4)) + ". Điền câu trả lời ngắn gọn vào phiếu làm bài.")
        p_p3_note.runs[0].font.italic = True
        p_p3_note.runs[0].font.size = Pt(10)

        for item in exam_content.get("phan_III_tra_loi_ngan", []):
            p_q = doc.add_paragraph()
            r_q = p_q.add_run(f"Câu {item['cau_so']}: ")
            r_q.font.bold = True
            p_q.add_run(item["cau_hoi"])

    # PHẦN IV / TỰ LUẬN
    doc.add_paragraph()
    p_p4 = doc.add_paragraph()
    r_p4 = p_p4.add_run(f"PHẦN { 'IV' if template_type == 'mau_2' else 'III' }. Tự luận ({summary.get('diem_tl', 3.0)} điểm)")
    r_p4.font.name = "Times New Roman"
    r_p4.font.size = Pt(11.5)
    r_p4.font.bold = True
    p_p4_note = doc.add_paragraph("Thí sinh trình bày chi tiết bài làm vào giấy kiểm tra.")
    p_p4_note.runs[0].font.italic = True
    p_p4_note.runs[0].font.size = Pt(10)

    for item in exam_content.get("phan_IV_tu_luan", []):
        p_q = doc.add_paragraph()
        r_q = p_q.add_run(f"Câu {item['cau_so']} ({item['diem']} điểm): ")
        r_q.font.bold = True
        p_q.add_run(item["cau_hoi"])

    doc.add_page_break()

    # ==============================================================
    # 4. HƯỚNG DẪN CHẤM VÀ ĐÁP ÁN CHI TIẾT
    # ==============================================================
    p_title_ans = doc.add_paragraph()
    r_ta = p_title_ans.add_run("IV. HƯỚNG DẪN CHẤM VÀ BIỂU ĐIỂM CHI TIẾT")
    r_ta.font.name = "Times New Roman"
    r_ta.font.size = Pt(13)
    r_ta.font.bold = True

    # 1. Đáp án trắc nghiệm nhiều lựa chọn
    p_ans_p1 = doc.add_paragraph()
    p_ans_p1.add_run("1. Đáp án Phần I (Trắc nghiệm nhiều lựa chọn - Mỗi câu đúng được 0,25 điểm)").bold = True

    ans_items = exam_content.get("phan_I_trac_nghiem", [])
    if ans_items:
        chunk_size = 12
        for start_idx in range(0, len(ans_items), chunk_size):
            chunk = ans_items[start_idx : start_idx + chunk_size]
            tbl_ans1 = doc.add_table(rows=2, cols=len(chunk) + 1)
            set_table_borders(tbl_ans1)
            set_table_margins(tbl_ans1, 40, 40, 50, 50)
            tbl_ans1.alignment = WD_TABLE_ALIGNMENT.CENTER

            cell_q = tbl_ans1.cell(0, 0)
            cell_q.text = "Câu"
            set_cell_background(cell_q, "EAECEE")
            cell_q.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p_q = cell_q.paragraphs[0]
            p_q.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_q.runs[0].font.name = "Times New Roman"
            p_q.runs[0].font.size = Pt(9.5)
            p_q.runs[0].font.bold = True

            cell_ans = tbl_ans1.cell(1, 0)
            cell_ans.text = "Đáp án"
            set_cell_background(cell_ans, "EAECEE")
            cell_ans.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p_ans = cell_ans.paragraphs[0]
            p_ans.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_ans.runs[0].font.name = "Times New Roman"
            p_ans.runs[0].font.size = Pt(9.5)
            p_ans.runs[0].font.bold = True

            for idx, it in enumerate(chunk):
                c_h = tbl_ans1.cell(0, idx + 1)
                c_a = tbl_ans1.cell(1, idx + 1)
                c_h.text = str(it["cau_so"])
                c_a.text = it["dap_an"]
                set_cell_background(c_h, "F8F9FA")
                c_h.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                c_a.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                
                p_h = c_h.paragraphs[0]
                p_h.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r_h = p_h.runs[0]
                r_h.font.name = "Times New Roman"
                r_h.font.size = Pt(9)
                r_h.font.bold = True

                p_a = c_a.paragraphs[0]
                p_a.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r_a = p_a.runs[0]
                r_a.font.name = "Times New Roman"
                r_a.font.size = Pt(10)
                r_a.font.bold = True

            doc.add_paragraph()

    # 2. Đáp án Phần II (Đúng/Sai)
    p_ans_p2 = doc.add_paragraph()
    p_ans_p2.add_run("2. Đáp án Phần II (Trắc nghiệm Đúng - Sai)").bold = True
    p_ds_note = doc.add_paragraph("Quy tắc chấm điểm chuẩn Bộ GD&ĐT:\n- Đúng 01 ý: 0,10 điểm\n- Đúng 02 ý: 0,25 điểm\n- Đúng 03 ý: 0,50 điểm\n- Đúng 04 ý: 1,00 điểm")
    p_ds_note.runs[0].font.italic = True
    p_ds_note.runs[0].font.size = Pt(9.5)

    ds_items = exam_content.get("phan_II_dung_sai", [])
    for it in ds_items:
        p_c = doc.add_paragraph()
        p_c.add_run(f"Câu {it['cau_so']}: ").bold = True
        for y in it.get("y_hoi", []):
            doc.add_paragraph(f"  • Ý {y['ky_hieu']}): {y['dap_an']} ({y.get('giai_thich', '')})")

    # 3. Đáp án Phần III (nếu là Mẫu 2)
    if template_type == "mau_2":
        doc.add_paragraph()
        p_ans_p3 = doc.add_paragraph()
        p_ans_p3.add_run("3. Đáp án Phần III (Trả lời ngắn - Mỗi câu đúng được 0,5 điểm)").bold = True
        for it in exam_content.get("phan_III_tra_loi_ngan", []):
            doc.add_paragraph(f"  • Câu {it['cau_so']}: Đáp án: {it['dap_an']} ({it.get('giai_thich', '')})")

    # 4. Biểu điểm Phần Tự luận
    doc.add_paragraph()
    p_ans_p4 = doc.add_paragraph()
    p_ans_p4.add_run("4. Biểu điểm và Hướng dẫn chấm Phần Tự luận").bold = True

    for it in exam_content.get("phan_IV_tu_luan", []):
        doc.add_paragraph(f"Câu {it['cau_so']} ({it['diem']} điểm): {it['cau_hoi']}")
        for step in it.get("huong_dan_cham", []):
            p_s = doc.add_paragraph()
            p_s.paragraph_format.left_indent = Inches(0.25)
            p_s.add_run(f"• {step['buoc']}: ").italic = True
            p_s.add_run(f"{step['diem']} điểm").bold = True

    # Lưu vào buffer bộ nhớ
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer
