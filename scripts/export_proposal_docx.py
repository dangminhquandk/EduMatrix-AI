import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_proposal_docx(md_path, docx_path):
    doc = Document()
    
    # Thiết lập lề chuẩn văn bản Nghị định 30/2020/NĐ-CP:
    # Trên: 20mm (0.79 in), Dưới: 20mm (0.79 in), Trái: 30mm (1.18 in), Phải: 15mm (0.59 in)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.79)
        section.bottom_margin = Inches(0.79)
        section.left_margin = Inches(1.18)
        section.right_margin = Inches(0.59)
    
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()
        
    for line in lines:
        line_str = line.strip()
        if not line_str:
            continue
            
        # Tiêu đề H1
        if line_str.startswith("# "):
            title_text = line_str[2:].strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(12)
            run = p.add_run(title_text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(16)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
            
        # Tiêu đề H2
        elif line_str.startswith("## "):
            h2_text = line_str[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(h2_text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(13.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0)
            
        # Tiêu đề H3
        elif line_str.startswith("### "):
            h3_text = line_str[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(h3_text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(12.5)
            run.font.bold = True
            run.font.italic = True
            run.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)
            
        # Dòng gạch ngang separator
        elif line_str.startswith("---"):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run("____________________________________________________________________")
            run.font.name = 'Times New Roman'
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0xCB, 0xD5, 0xE0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
        # Danh sách gạch đầu dòng / số thứ tự
        elif re.match(r"^(\*|\-|\d+\.)\s+", line_str):
            clean_text = re.sub(r"^(\*|\-|\d+\.)\s+", "", line_str)
            p = doc.add_paragraph(style='List Bullet' if line_str.startswith(('*', '-')) else 'List Number')
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.2
            
            # Xử lý bold trong dòng
            parts = re.split(r'(\*\*.*?\*\*)', clean_text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.font.name = 'Times New Roman'
                    run.font.size = Pt(12)
                    run.font.bold = True
                else:
                    run = p.add_run(part)
                    run.font.name = 'Times New Roman'
                    run.font.size = Pt(12)
                    
        # Đoạn văn thông thường
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.25
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            
            # Xử lý bold trong đoạn văn
            parts = re.split(r'(\*\*.*?\*\*)', line_str)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.font.name = 'Times New Roman'
                    run.font.size = Pt(12)
                    run.font.bold = True
                else:
                    run = p.add_run(part)
                    run.font.name = 'Times New Roman'
                    run.font.size = Pt(12)
                    
    doc.save(docx_path)
    print(f"File created successfully: {docx_path}")

if __name__ == '__main__':
    md = "/Users/quandkh/Documents/Research_Hub/Sản phẩm số AI/Ban_Thuyet_Minh_Sang_Kien.md"
    out = "/Users/quandkh/Documents/Research_Hub/Sản phẩm số AI/Ban_Thuyet_Minh_Sang_Kien.docx"
    create_proposal_docx(md, out)
